"""Primary GOLD joins preserve source chains without guessing genome identity."""

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from taxonmech.gold_genomes import _HEADERS, extract_gold_genomes

STRAINS = [
    {"strain_id": "kgmicrobe.strain:bacdive_1", "culture_collection_ids": "kgmicrobe.strain:DSM-6724"},
    {"strain_id": "kgmicrobe.strain:bacdive_2", "culture_collection_ids": "kgmicrobe.strain:ATCC-BAA-1556"},
]


def organism(identifier="Go1", culture="DSM 6724", strain="", **extra):
    return {"ORGANISM GOLD ID": identifier, "ORGANISM NAME": "Source organism name",
            "ORGANISM NCBI TAX ID": "515635", "ORGANISM STRAIN": strain,
            "ORGANISM CULTURE COLLECTION ID": culture, **extra}


def project(identifier="Gp1", oid="Go1", **extra):
    return {"PROJECT GOLD ID": identifier, "PROJECT NAME": "Source project name",
            "ORGANISM GOLD ID": oid, "NCBI BIOPROJECT ACCESSION": "PRJNA29175",
            "NCBI BIOSAMPLE ACCESSION": "SAMN00002661", **extra}


def analysis(identifier="Ga1", projects="Gp1", oid="", **extra):
    return {"AP GOLD ID": identifier, "AP NAME": "Source analysis name",
            "AP TYPE": "Genome Analysis (Isolate)",
            "AP IMG TAXON ID": "2517572146", "AP ORGANISM GOLD ID": oid, "AP PROJECT GOLD IDS": projects,
            "AP GENBANK": json.dumps([{"genbankId": "AQXM00000000", "assemblyAccession": "GCA_000376245.1"}]),
            **extra}


def workbook(tmp_path: Path, organisms=None, projects=None, analyses=None, bad_dimensions=False):
    path = tmp_path / "goldData.xlsx"
    book = Workbook()
    book.remove(book.active)
    for sheet, rows in (
        ("Organism", [organism()] if organisms is None else organisms),
        ("Sequencing Project", [project()] if projects is None else projects),
        ("Analysis Project", [analysis()] if analyses is None else analyses),
    ):
        ws = book.create_sheet(sheet)
        ws.append(_HEADERS[sheet])
        for row in rows:
            ws.append([row.get(field, "") for field in _HEADERS[sheet]])
    book.save(path)
    if bad_dimensions:
        import re

        replacement = tmp_path / "rewritten.xlsx"
        with ZipFile(path) as source, ZipFile(replacement, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename.startswith("xl/worksheets/"):
                    payload = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', payload)
                target.writestr(info, payload)
        replacement.replace(path)
    return path


def test_primary_workbook_chain_and_incorrect_declared_dimensions(tmp_path):
    assemblies, genomes, related, drops = extract_gold_genomes(
        workbook(tmp_path, bad_dimensions=True), STRAINS,
    )
    assert not drops
    assert [row["assembly_id"] for row in assemblies] == ["ncbi.assembly:GCA_000376245.1"]
    assert [row["genome_id"] for row in genomes] == ["img.taxon:2517572146"]
    assert {row["record_type"] for row in related} == {
        "GOLD_ORGANISM", "GOLD_PROJECT", "GOLD_ANALYSIS", "BIOSAMPLE", "BIOPROJECT",
    }
    assert all(row["source"] == "GOLD" for row in [*assemblies, *genomes, *related])
    evidence = assemblies[0]
    assert evidence["source_id"] == "gold:Ga1"
    assert evidence["source_organism_id"] == "gold:Go1"
    assert evidence["source_project_id"] == "gold:Gp1"
    assert evidence["source_field"] == "AP GENBANK.assemblyAccession"
    assert evidence["matched_strain_id"] == "kgmicrobe.strain:DSM-6724"
    assert evidence["source_strain_field"] == "ORGANISM CULTURE COLLECTION ID"
    assert evidence["source_strain_identifiers"] == "DSM 6724"
    assert {row["record_name"] for row in related if row["record_type"].startswith("GOLD_")} == {
        "Source organism name", "Source project name", "Source analysis name",
    }


@pytest.mark.parametrize("culture", ["6724", "DSM 06724", "DSM 6724T", "Genus species DSM 6724"])
def test_taxonomy_names_numbers_and_substrings_never_join(tmp_path, culture):
    result = extract_gold_genomes(workbook(tmp_path, organisms=[organism(culture=culture)]), STRAINS)
    assert result == ([], [], [], [])


def test_whole_deposits_in_strain_column_and_multiple_fields_retain_evidence(tmp_path):
    raw = "other designation, DSM 6724; ATCC BAA-1556"
    assemblies, _, _, drops = extract_gold_genomes(
        workbook(tmp_path, organisms=[organism(strain=raw)]), STRAINS,
    )
    assert not drops
    assert len(assemblies) == 3
    assert {row["strain_id"] for row in assemblies} == {row["strain_id"] for row in STRAINS}
    assert {row["source_strain_field"] for row in assemblies} == {
        "ORGANISM STRAIN", "ORGANISM CULTURE COLLECTION ID",
    }
    assert {row["source_strain_identifiers"] for row in assemblies} == {raw, "DSM 6724"}


@pytest.mark.parametrize("second_oid,direct_oid", [("Go2", ""), ("", ""), ("Go1", "Go2")])
def test_ambiguous_analysis_does_not_assign_every_genome_to_every_strain(tmp_path, second_oid, direct_oid):
    path = workbook(
        tmp_path, organisms=[organism(), organism("Go2", "ATCC BAA-1556")],
        projects=[project(), project("Gp2", second_oid)],
        analyses=[analysis(projects="Gp1|Gp2", oid=direct_oid)],
    )
    assemblies, genomes, related, drops = extract_gold_genomes(path, STRAINS)
    assert not assemblies and not genomes
    assert all(row["record_type"] != "GOLD_ANALYSIS" for row in related)
    assert any("unambiguously" in row["reason"] for row in drops)


def test_missing_project_prevents_partial_chain_join(tmp_path):
    path = workbook(tmp_path, analyses=[analysis(projects="Gp1|Gp999")])
    assemblies, genomes, _, drops = extract_gold_genomes(path, STRAINS)
    assert not assemblies and not genomes and drops


def test_multiple_projects_for_same_organism_preserve_each_chain(tmp_path):
    path = workbook(tmp_path, projects=[project(), project("Gp2")], analyses=[analysis(projects="Gp1|Gp2")])
    assemblies, genomes, _, drops = extract_gold_genomes(path, STRAINS)
    assert not drops
    assert len(assemblies) == len(genomes) == 2
    assert {row["source_project_id"] for row in genomes} == {"gold:Gp1", "gold:Gp2"}


def test_direct_organism_reference_without_projects(tmp_path):
    path = workbook(tmp_path, projects=[], analyses=[analysis(projects="", oid="Go1")])
    assemblies, genomes, _, drops = extract_gold_genomes(path, STRAINS)
    assert not drops
    assert len(assemblies) == len(genomes) == 1
    assert assemblies[0]["source_organism_id"] == "gold:Go1"
    assert assemblies[0]["source_project_id"] == ""


def test_sequence_accessions_are_not_genomes_and_versions_are_preserved(tmp_path):
    genbank = [
        {"genbankId": "AQXM00000000"},
        {"genbankId": "CP000001", "assemblyAccession": "GCA_000376245"},
        {"assemblyAccession": "GCF_000376245.23"},
        {"assemblyAccession": "GCA_000376245.23"},
        {"assemblyAccession": "GCA_000376245.23"},
    ]
    path = workbook(tmp_path, analyses=[analysis(**{"AP GENBANK": json.dumps(genbank)})])
    assemblies, _, _, drops = extract_gold_genomes(path, STRAINS)
    assert not drops
    assert {row["assembly_id"] for row in assemblies} == {
        "ncbi.assembly:GCA_000376245", "ncbi.assembly:GCF_000376245.23", "ncbi.assembly:GCA_000376245.23",
    }
    assert len(assemblies) == 3


@pytest.mark.parametrize("field,value", [
    ("AP IMG TAXON ID", "GCF_000376245.1"),
    ("AP GENBANK", "not JSON"),
    ("AP GENBANK", '{"assemblyAccession":"GCA_000376245.1"}'),
    ("AP GENBANK", '[{"assemblyAccession":"CP000001"}]'),
])
def test_malformed_links_are_quarantined_without_erasing_valid_counterpart(tmp_path, field, value):
    path = workbook(tmp_path, analyses=[analysis(**{field: value})])
    assemblies, genomes, related, drops = extract_gold_genomes(path, STRAINS)
    assert drops and related
    if field == "AP IMG TAXON ID":
        assert assemblies and not genomes
    else:
        assert genomes and not assemblies


def test_transcriptome_analysis_stays_related_and_is_not_a_genome(tmp_path):
    path = workbook(tmp_path, analyses=[analysis(**{"AP TYPE": "Transcriptome Analysis"})])
    assemblies, genomes, related, drops = extract_gold_genomes(path, STRAINS)
    assert not assemblies and not genomes
    assert any(row["record_type"] == "GOLD_ANALYSIS" for row in related)
    assert any("AP TYPE" in row["reason"] for row in drops)


def test_malformed_related_id_does_not_remove_other_links(tmp_path):
    path = workbook(tmp_path, projects=[project(**{"NCBI BIOSAMPLE ACCESSION": "GCA_000376245.1"})])
    assemblies, genomes, related, drops = extract_gold_genomes(path, STRAINS)
    assert assemblies and genomes and drops
    assert all(row["record_type"] != "BIOSAMPLE" for row in related)
    assert any(row["record_type"] == "BIOPROJECT" for row in related)


def test_duplicate_source_identity_is_structural_error(tmp_path):
    with pytest.raises(ValueError, match="duplicate ID"):
        extract_gold_genomes(workbook(tmp_path, organisms=[organism(), organism()]), STRAINS)


def test_missing_source_sheet_is_structural_error(tmp_path):
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    with pytest.raises(ValueError, match="missing sheet"):
        extract_gold_genomes(path, STRAINS)

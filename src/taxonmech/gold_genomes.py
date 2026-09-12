"""Strain-linked genome and related records from GOLD's public workbook.

Organism, sequencing project and analysis project IDs retain distinct types.
Only explicit culture identifiers join to an inventoried strain. Analysis
genomes must resolve to one organism through every stated project; taxonomy
and organism names never establish a strain or genome identity.
"""

from __future__ import annotations

import json
import re
from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

from taxonmech.genome_sources import index_culture_identifiers, normalize_culture_identifier

_HEADERS = {
    "Organism": (
        "ORGANISM GOLD ID", "ORGANISM NAME", "ORGANISM NCBI TAX ID",
        "ORGANISM STRAIN", "ORGANISM CULTURE COLLECTION ID",
    ),
    "Sequencing Project": (
        "PROJECT GOLD ID", "PROJECT NAME", "ORGANISM GOLD ID",
        "NCBI BIOPROJECT ACCESSION", "NCBI BIOSAMPLE ACCESSION",
    ),
    "Analysis Project": (
        "AP GOLD ID", "AP NAME", "AP TYPE", "AP IMG TAXON ID", "AP GENBANK",
        "AP ORGANISM GOLD ID", "AP PROJECT GOLD IDS",
    ),
}
_GENOME_ANALYSES = {
    "Genome Analysis (Isolate)", "Combined Assembly", "Metagenome-Assembled Genome",
    "Single Cell Analysis (screened)", "Single Cell Analysis (unscreened)",
    "Combined Assembly Single Cell (screened)", "Combined Assembly Single Cell (unscreened)",
    "Metagenome-Extracted Genome",
}
_ASSEMBLY = re.compile(r"GC[AF]_[0-9]{9}(?:\.[1-9][0-9]*)?")
_BIOSAMPLE = re.compile(r"SAM[EDN][A-Z]?[0-9]+")
_BIOPROJECT = re.compile(r"PRJ(?:NA|DB|DA|EB|EA)[0-9]+")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _rows(workbook, sheet: str) -> Iterator[dict[str, str]]:
    if sheet not in workbook.sheetnames:
        raise ValueError(f"GOLD workbook is missing sheet {sheet!r}")
    worksheet = workbook[sheet]
    # GOLD's public export declares A1:A1 even for sheets with 500,000 rows.
    # Trusting that dimension silently discards every data row in read-only mode.
    worksheet.reset_dimensions()
    iterator = worksheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(iterator, ())]
    missing = set(_HEADERS[sheet]) - set(headers)
    if missing or len(headers) != len(set(headers)):
        raise ValueError(f"GOLD {sheet}: missing or duplicate headers: {sorted(missing)}")
    positions = {field: headers.index(field) for field in _HEADERS[sheet]}
    for number, values in enumerate(iterator, 2):
        if not any(value is not None for value in values):
            continue
        if len(values) > len(headers):
            raise ValueError(f"GOLD {sheet} row {number}: more cells than headers")
        yield {field: _text(values[index]) if index < len(values) else ""
               for field, index in positions.items()}


def _tokens(value: str) -> list[str]:
    return [token.strip() for token in re.split(r"[|;,]", value) if token.strip()]


def _gold_id(value: str, kind: str) -> bool:
    return re.fullmatch(rf"G{kind}[0-9]+", value) is not None


def _unique(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    found = {tuple(sorted(row.items())): row for row in rows}
    return [found[key] for key in sorted(found)]


def extract_gold_genomes(
    path: Path, strain_rows: list[dict[str, str]],
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    """Return assemblies, IMG records, typed related records and exclusions.

    Every asserted chain retains its GOLD organism and sequencing project IDs,
    the exact inventoried deposit, and the source strain column and text.
    GOLD Go/Gp/Ga IDs are related records, never mislabeled as genome records.
    """
    deposits = index_culture_identifiers(strain_rows)
    assemblies: list[dict[str, str]] = []
    genomes: list[dict[str, str]] = []
    related: list[dict[str, str]] = []
    drops: list[dict[str, str]] = []
    organisms: dict[str, list[dict[str, str]]] = {}
    organism_ids: set[str] = set()
    projects: dict[str, str] = {}

    def reject(kind: str, identifier: str, reason: str) -> None:
        drops.append({"kind": f"gold_{kind}", "id": identifier, "reason": reason})

    def related_row(provenance: dict[str, str], identifier: str, kind: str, field: str, name: str) -> None:
        related.append({**provenance, "record_id": identifier, "record_type": kind,
                        "source_field": field, "record_name": name})

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for row in _rows(workbook, "Organism"):
            oid = row["ORGANISM GOLD ID"].strip()
            if not _gold_id(oid, "o") or oid in organism_ids:
                raise ValueError(f"GOLD Organism: invalid or duplicate ID {oid!r}")
            organism_ids.add(oid)
            matches = []
            for field in ("ORGANISM CULTURE COLLECTION ID", "ORGANISM STRAIN"):
                for token in _tokens(row[field]):
                    for strain_id, deposit in sorted(deposits.get(normalize_culture_identifier(token), ())):
                        matches.append({
                            "strain_id": strain_id, "matched_strain_id": deposit,
                            "source_strain_field": field, "source_strain_identifiers": row[field],
                        })
            if not matches:
                continue
            taxon = row["ORGANISM NCBI TAX ID"].strip()
            if taxon and not re.fullmatch(r"[1-9][0-9]*", taxon):
                reject("organism", f"gold:{oid}", "malformed ORGANISM NCBI TAX ID")
                continue
            provenance = {
                "source": "GOLD", "source_id": f"gold:{oid}", "source_organism_id": f"gold:{oid}",
                "source_project_id": "", "source_reference_id": "",
                "taxon_id": f"NCBITaxon:{taxon}" if taxon else "",
            }
            organisms[oid] = _unique([{**provenance, **match} for match in matches])
            for evidence in organisms[oid]:
                related_row(evidence, f"gold:{oid}", "GOLD_ORGANISM", "ORGANISM GOLD ID",
                            row["ORGANISM NAME"])

        for row in _rows(workbook, "Sequencing Project"):
            pid = row["PROJECT GOLD ID"].strip()
            if not _gold_id(pid, "p") or pid in projects:
                raise ValueError(f"GOLD Sequencing Project: invalid or duplicate ID {pid!r}")
            oid = row["ORGANISM GOLD ID"].strip()
            projects[pid] = oid
            for evidence in organisms.get(oid, []):
                provenance = {**evidence, "source_id": f"gold:{pid}", "source_project_id": f"gold:{pid}"}
                related_row(provenance, f"gold:{pid}", "GOLD_PROJECT", "PROJECT GOLD ID", row["PROJECT NAME"])
                for field, prefix, kind, pattern in (
                    ("NCBI BIOSAMPLE ACCESSION", "biosample", "BIOSAMPLE", _BIOSAMPLE),
                    ("NCBI BIOPROJECT ACCESSION", "bioproject", "BIOPROJECT", _BIOPROJECT),
                ):
                    for accession in _tokens(row[field]):
                        if not pattern.fullmatch(accession):
                            reject("related_record", f"gold:{pid}/{accession}", f"malformed {field}")
                            continue
                        related_row(provenance, f"{prefix}:{accession}", kind, field, "")

        analysis_ids = set()
        for row in _rows(workbook, "Analysis Project"):
            aid = row["AP GOLD ID"].strip()
            if not _gold_id(aid, "a") or aid in analysis_ids:
                raise ValueError(f"GOLD Analysis Project: invalid or duplicate ID {aid!r}")
            analysis_ids.add(aid)
            pids = sorted(set(_tokens(row["AP PROJECT GOLD IDS"])))
            direct_oid = row["AP ORGANISM GOLD ID"].strip()
            chain_oids = {projects.get(pid, "") for pid in pids}
            if direct_oid:
                chain_oids.add(direct_oid)
            # Do not spend exclusions on unrelated portions of the public database.
            if not any(oid in organisms for oid in chain_oids):
                continue
            if (len(chain_oids) != 1 or "" in chain_oids
                    or any(not _gold_id(pid, "p") or pid not in projects for pid in pids)
                    or any(oid not in organism_ids for oid in chain_oids)):
                reject("analysis", f"gold:{aid}", "analysis does not resolve unambiguously to one organism")
                continue
            oid = next(iter(chain_oids))
            chains = [{**evidence, "source_id": f"gold:{aid}",
                       "source_project_id": f"gold:{pid}" if pid else ""}
                      for pid in (pids or [""]) for evidence in organisms[oid]]
            for evidence in chains:
                related_row(evidence, f"gold:{aid}", "GOLD_ANALYSIS", "AP GOLD ID", row["AP NAME"])
            if row["AP TYPE"].strip() not in _GENOME_ANALYSES:
                if row["AP IMG TAXON ID"].strip() or row["AP GENBANK"].strip():
                    reject("analysis", f"gold:{aid}", f"non-genome or unsupported AP TYPE: {row['AP TYPE']}")
                continue
            img_id = row["AP IMG TAXON ID"].strip()
            if img_id and not re.fullmatch(r"[1-9][0-9]*", img_id):
                reject("genome_record", f"gold:{aid}/{img_id}", "malformed AP IMG TAXON ID")
                img_id = ""
            targets = []
            raw_genbank = row["AP GENBANK"].strip()
            if raw_genbank:
                try:
                    records = json.loads(raw_genbank)
                except (ValueError, TypeError):
                    records = None
                if not isinstance(records, list) or any(not isinstance(item, dict) for item in records):
                    reject("assembly", f"gold:{aid}", "malformed AP GENBANK JSON array")
                else:
                    for item in records:
                        accession = item.get("assemblyAccession")
                        if accession in (None, ""):
                            continue
                        if not isinstance(accession, str) or not _ASSEMBLY.fullmatch(accession.strip()):
                            reject("assembly", f"gold:{aid}/{accession}",
                                   "malformed AP GENBANK.assemblyAccession")
                            continue
                        targets.append(accession.strip())
            for evidence in chains:
                if img_id:
                    genomes.append({**evidence, "genome_id": f"img.taxon:{img_id}", "source_database": "img",
                                    "source_field": "AP IMG TAXON ID", "assembly_level": "",
                                    "genome_name": row["AP NAME"]})
                for accession in targets:
                    assemblies.append({**evidence, "assembly_id": f"ncbi.assembly:{accession}",
                                       "source_field": "AP GENBANK.assemblyAccession", "assembly_level": "",
                                       "assembly_name": row["AP NAME"]})
    finally:
        workbook.close()
    return tuple(_unique(rows) for rows in (assemblies, genomes, related, drops))
